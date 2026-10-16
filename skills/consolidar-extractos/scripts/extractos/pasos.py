"""Los pasos del flujo. Cada uno lee/escribe archivos CSV en <salida>/trabajo/ y no toca los originales.

  1 inventario  -> inventario.csv       que hay, de que banco/cuenta, que rango, duplicados, PDFs sin hermano
  2 leer        -> crudos.csv           todos los xls/xlsx/csv al esquema comun (con solapes)
  3 unificar    -> unificados.csv       un solo archivo por dia y cuenta; fuentes.csv; cortes.csv
  4 pdf         -> final.csv            suma SOLO las filas de PDFs sin version excel que no estan ya
  4b mayor      -> mayor.csv            (opcional) tramos sin extracto, desde el libro mayor, en hojas aparte
  5 excel       -> <Banco> ... .xlsx    un Excel por banco + interno/Trazabilidad
  6 verificar   -> informe              recalcula todo desde los Excel
"""
import collections
import datetime
import json
import os
import re
from decimal import Decimal

from .adaptadores import leer_archivo, leer_mayor
from .comunes import (CERO, COLUMNAS, TOL, a_numero, errores_de_saldo, escribir_movimientos, escribir_tabla,
                      leer_movimientos, md5)

EXT_TABLA = ('xls', 'xlsx', 'csv', 'txt')
EXT_TODAS = EXT_TABLA + ('pdf',)


def _trabajo(cfg):
    os.makedirs(cfg['trabajo'], exist_ok=True)
    return cfg['trabajo']


def _archivos(cfg):
    for dp, _, fs in os.walk(cfg['entrada']):
        for fn in sorted(fs):
            if fn.lower().rsplit('.', 1)[-1] in EXT_TODAS and not fn.startswith('~$'):
                p = os.path.join(dp, fn)
                yield p, os.path.relpath(p, cfg['entrada'])


def _tallo(nombre):
    return re.sub(r'^\d+\.\s*', '', nombre.rsplit('.', 1)[0]).strip().lower()


# ------------------------------------------------------------------ 1. inventario
def inventario(cfg):
    lista = list(_archivos(cfg))
    por_md5 = collections.defaultdict(list)
    hermanos = collections.defaultdict(set)
    infos = []
    for p, rel in lista:
        h = md5(p)
        ext = rel.lower().rsplit('.', 1)[-1]
        por_md5[h].append(rel)
        hermanos[(os.path.dirname(rel), _tallo(os.path.basename(rel)))].add(ext)
        infos.append((p, rel, ext, h))
    filas = []
    for p, rel, ext, h in infos:
        marcas = []
        if len(por_md5[h]) > 1:
            marcas.append('IDENTICO a otro archivo')
        if ext == 'pdf':
            if hermanos[(os.path.dirname(rel), _tallo(os.path.basename(rel)))] & set(EXT_TABLA):
                marcas.append('PDF con version excel/csv')
            else:
                marcas.append('PDF sin hermano')
        info = dict(archivo=rel, ext=ext, banco='', cuenta='', formato='', filas='', desde='', hasta='',
                    cortes_saldo='', marcas='; '.join(marcas), avisos='', md5=h)
        try:
            r = leer_archivo(p, rel, cfg) if ext in EXT_TABLA else None
            if r is None and ext == 'pdf':
                from .adaptadores import detectar_banco, asignar_cuenta
                info['banco'] = detectar_banco(rel, '', cfg)
                info['cuenta'], _ = asignar_cuenta(info['banco'], rel, '', cfg)
            elif r:
                info.update(banco=r['banco'], cuenta=r['cuenta'], formato=r['formato'], filas=len(r['filas']),
                            avisos=' / '.join(r['avisos']))
                if r['filas']:
                    info.update(desde=min(f['fecha'] for f in r['filas']), hasta=max(f['fecha'] for f in r['filas']),
                                cortes_saldo=errores_de_saldo(r['filas']))
        except Exception as e:      # un archivo roto no frena el inventario: se informa
            info['avisos'] = f'ERROR al leer: {type(e).__name__}: {e}'
        filas.append(info)
    filas.sort(key=lambda f: (f['banco'], f['cuenta'], str(f['desde']), f['archivo']))
    cols = ['banco', 'cuenta', 'archivo', 'ext', 'formato', 'filas', 'desde', 'hasta', 'cortes_saldo', 'marcas', 'avisos', 'md5']
    escribir_tabla(os.path.join(_trabajo(cfg), 'inventario.csv'), filas, cols)
    res = collections.Counter(f['ext'] for f in filas)
    print(f'{len(filas)} archivos: ' + ', '.join(f'{v} {k}' for k, v in sorted(res.items())))
    cuentas = collections.defaultdict(list)
    for f in filas:
        cuentas[(f['banco'], f['cuenta'])].append(f)
    for (b, c), fs in sorted(cuentas.items()):
        d = [str(x['desde']) for x in fs if x['desde']]
        h = [str(x['hasta']) for x in fs if x['hasta']]
        print(f'  {b:22} {c:34} {len(fs):3} archivos  {min(d) if d else "-"} a {max(h) if h else "-"}')
    sin = [f for f in filas if f['cuenta'] in ('(sin identificar)', '') and f['ext'] != 'pdf']
    desconocidos = [f for f in filas if f['banco'] == 'Banco sin identificar' or f['formato'] == 'no reconocido']
    if sin:
        print(f'\nATENCION: {len(sin)} archivos sin cuenta identificada: preguntar al usuario y agregar reglas en "cuentas".')
    if desconocidos:
        print(f'ATENCION: {len(desconocidos)} archivos con banco o formato no reconocido (ver inventario.csv).')
    return filas


# ------------------------------------------------------------------ 2. leer
def leer(cfg):
    movs, control, solapas = [], [], {}
    for p, rel in _archivos(cfg):
        if rel.lower().rsplit('.', 1)[-1] not in EXT_TABLA:
            continue
        r = leer_archivo(p, rel, cfg)
        if r is None or not r['filas']:
            control.append(dict(archivo=rel, banco=r['banco'] if r else '', cuenta='', filas=0,
                                cortes_saldo='', nota='; '.join(r['avisos']) if r else 'sin filas'))
            continue
        if r['solapa']:
            solapas[f"{r['banco']}||{r['cuenta']}"] = r['solapa']
        for f in r['filas']:
            movs.append(dict(f, banco=r['banco'], cuenta=r['cuenta'], archivo=rel, nota=''))
        control.append(dict(archivo=rel, banco=r['banco'], cuenta=r['cuenta'], filas=len(r['filas']),
                            desde=r['filas'][0]['fecha'], hasta=r['filas'][-1]['fecha'],
                            cortes_saldo=errores_de_saldo(r['filas']), nota='; '.join(r['avisos'])))
    t = _trabajo(cfg)
    escribir_movimientos(os.path.join(t, 'crudos.csv'), movs)
    escribir_tabla(os.path.join(t, 'control_lectura.csv'), control,
                   ['archivo', 'banco', 'cuenta', 'filas', 'desde', 'hasta', 'cortes_saldo', 'nota'])
    json.dump(solapas, open(os.path.join(t, 'solapas.json'), 'w', encoding='utf-8'), ensure_ascii=False)
    malos = [c for c in control if c.get('cortes_saldo') not in ('', 0)]
    print(f'{len(control)} archivos leidos, {len(movs)} movimientos (con solapes aun sin quitar)')
    print(f'Archivos con saldo corrido que no cierra: {len(malos)}' + (' -> ver control_lectura.csv' if malos else ''))
    return movs


# ------------------------------------------------------------------ 3. unificar
def _cortes(banco, cuenta, cadena):
    salida = []
    for a, b in zip(cadena, cadena[1:]):
        if a['saldo'] is None or b['saldo'] is None:
            continue
        if abs(a['saldo'] + b['credito'] - b['debito'] - b['saldo']) > TOL:
            salida.append(dict(banco=banco, cuenta=cuenta, fecha=b['fecha'], descripcion=b['descripcion'][:60],
                               saldo_anterior=a['saldo'], saldo=b['saldo'],
                               diferencia=b['saldo'] - (a['saldo'] + b['credito'] - b['debito']), archivo=b['archivo']))
    return salida


def _por_cuenta(movs):
    grupos = collections.OrderedDict()
    for r in movs:
        grupos.setdefault((r['banco'], r['cuenta']), collections.OrderedDict()).setdefault(r['archivo'], []).append(r)
    return grupos


def _encadenar(saldo, filas):
    """Ordena las filas de un dia para que cada saldo salga del anterior. None si no se puede (no se fuerza)."""
    restantes, salida = list(filas), []
    while restantes:
        for i, r in enumerate(restantes):
            if r['saldo'] is not None and abs(saldo + r['credito'] - r['debito'] - r['saldo']) <= TOL:
                salida.append(restantes.pop(i))
                saldo = r['saldo']
                break
        else:
            return None
    return salida


def unificar(cfg):
    t = _trabajo(cfg)
    movs = leer_movimientos(os.path.join(t, 'crudos.csv'))
    final, fuentes, cortes = [], [], []
    # Regla: por dia y cuenta se usa UN solo archivo (el que tiene saldo, luego mas filas ese dia, luego el
    # mas grande). Lo que otro archivo tenga de mas ese dia se agrega marcado. Las filas se comparan por
    # fecha + debito + credito, NUNCA por texto: el banco cambia descripciones entre descargas.
    for (banco, cuenta), archivos in sorted(_por_cuenta(movs).items()):
        tam = {a: len(fs) for a, fs in archivos.items()}
        por_dia = collections.defaultdict(list)
        for a, fs in archivos.items():
            dias = collections.OrderedDict()
            for r in fs:
                dias.setdefault(r['fecha'], []).append(r)
            for d, filas_dia in dias.items():
                por_dia[d].append((a, filas_dia))
        cadena, usadas = [], collections.Counter()
        clave = lambda r: (r['fecha'], r['debito'], r['credito'])
        for d in sorted(por_dia):
            cand = sorted(por_dia[d], key=lambda x: (x[1][0]['saldo'] is None, -len(x[1]), -tam[x[0]]))
            a0, base = cand[0]
            usadas[a0] += len(base)
            tiene = collections.Counter(clave(r) for r in base)
            extras = []
            for a, filas_dia in cand[1:]:
                falta = collections.Counter(clave(r) for r in filas_dia) - tiene
                for r in filas_dia:
                    if falta[clave(r)] > 0:
                        falta[clave(r)] -= 1
                        tiene[clave(r)] += 1
                        extras.append(dict(r, nota=f'solo en archivo: {os.path.basename(a)}'))
                        usadas[a] += 1
            dia = list(base) + extras
            if extras and cadena and cadena[-1]['saldo'] is not None:
                # dia partido entre dos archivos: se ordena siguiendo el saldo, no el orden de los archivos
                ordenado = _encadenar(cadena[-1]['saldo'], dia)
                if ordenado is not None:
                    dia = ordenado
            cadena.extend(dia)
        for a, fs in archivos.items():
            u = usadas[a]
            estado = ('descartado (todo su contenido ya estaba)' if u == 0 else 'usado completo' if u == len(fs)
                      else f'parcial ({len(fs) - u} filas no usadas)')
            fuentes.append(dict(banco=banco, cuenta=cuenta, archivo=a, desde=fs[0]['fecha'], hasta=fs[-1]['fecha'],
                                filas_archivo=len(fs), filas_usadas=u, estado=estado))
        cortes.extend(_cortes(banco, cuenta, cadena))
        final.extend(cadena)
    escribir_movimientos(os.path.join(t, 'unificados.csv'), final)
    escribir_tabla(os.path.join(t, 'fuentes.csv'), fuentes,
                   ['banco', 'cuenta', 'archivo', 'desde', 'hasta', 'filas_archivo', 'filas_usadas', 'estado'])
    escribir_tabla(os.path.join(t, 'cortes.csv'), cortes,
                   ['banco', 'cuenta', 'fecha', 'descripcion', 'saldo_anterior', 'saldo', 'diferencia', 'archivo'])
    print(f'Crudos {len(movs)} -> unificados {len(final)} (quitados {len(movs) - len(final)} repetidos)')
    _resumen_cuentas(final, cortes)
    return final


def _resumen_cuentas(filas, cortes):
    por = collections.OrderedDict()
    for r in filas:
        por.setdefault((r['banco'], r['cuenta']), []).append(r)
    for (b, c), fs in por.items():
        n = sum(1 for x in cortes if (x['banco'], x['cuenta']) == (b, c))
        print(f'  {b:22} {c:34} {len(fs):5} filas  {fs[0]["fecha"]} a {fs[-1]["fecha"]}  cortes de saldo: {n}')


# ------------------------------------------------------------------ 4. pdf
def pdf(cfg):
    t = _trabajo(cfg)
    base = leer_movimientos(os.path.join(t, 'unificados.csv'))
    inv = list(__import__('csv').DictReader(open(os.path.join(t, 'inventario.csv'), encoding='utf-8-sig'), delimiter=';'))
    sin_hermano = [f for f in inv if f['ext'] == 'pdf' and 'PDF sin hermano' in f['marcas']]
    clave_f = lambda r: (r['banco'], r['cuenta'], r['fecha'], r['debito'], r['credito'], r['saldo'])
    contador = collections.Counter(clave_f(r) for r in base)
    nuevas, control, vistos_md5 = collections.defaultdict(list), [], set()
    for f in sin_hermano:
        if f['md5'] in vistos_md5:
            control.append(dict(archivo=f['archivo'], filas='', nuevas='', nota='idéntico a otro PDF ya leído'))
            continue
        vistos_md5.add(f['md5'])
        p = os.path.join(cfg['entrada'], f['archivo'])
        r = leer_archivo(p, f['archivo'], cfg)
        filas = [dict(x, banco=r['banco'], cuenta=r['cuenta'], archivo=f['archivo']) for x in r['filas']]
        tiene = collections.Counter(contador)            # cada PDF se compara por separado contra lo unificado
        agregadas = []
        for x in filas:
            k = clave_f(x)
            if tiene[k] > 0:
                tiene[k] -= 1
            else:
                agregadas.append(x)
        # entre PDFs distintos que traen la misma fila nueva: se queda una
        for x in agregadas:
            nuevas[(r['banco'], r['cuenta'])].append(x)
        control.append(dict(archivo=f['archivo'], filas=len(filas), nuevas=len(agregadas),
                            nota=('; '.join(r['avisos']) + ' ') + f'saldo del PDF: {errores_de_saldo(r["filas"])} cortes'))
    final, cortes, sumadas = [], [], 0
    grupos = collections.OrderedDict()
    for r in base:
        grupos.setdefault((r['banco'], r['cuenta']), []).append(r)
    for k in nuevas:
        grupos.setdefault(k, [])
    for k, filas in grupos.items():
        extra, tope, usados_pdf = [], collections.Counter(), collections.defaultdict(collections.Counter)
        por_pdf = collections.OrderedDict()
        for x in nuevas.get(k, []):
            por_pdf.setdefault(x['archivo'], []).append(x)
        for arch, rs in por_pdf.items():
            cnt = collections.Counter(clave_f(x) for x in rs)
            usados = collections.Counter()
            for x in rs:
                c = clave_f(x)
                usados[c] += 1
                if usados[c] > tope[c]:
                    extra.append(dict(x, nota=f'solo en PDF: {os.path.basename(arch)}'))
            for c, v in cnt.items():
                tope[c] = max(tope[c], v)
        extra.sort(key=lambda x: x['fecha'])
        cadena = filas + extra
        sumadas += len(extra)
        cortes.extend(_cortes(k[0], k[1], cadena))
        final.extend(cadena)
    escribir_movimientos(os.path.join(t, 'final.csv'), final)
    escribir_tabla(os.path.join(t, 'cortes_final.csv'), cortes,
                   ['banco', 'cuenta', 'fecha', 'descripcion', 'saldo_anterior', 'saldo', 'diferencia', 'archivo'])
    escribir_tabla(os.path.join(t, 'control_pdf.csv'), control, ['archivo', 'filas', 'nuevas', 'nota'])
    print(f'{len(sin_hermano)} PDFs sin version excel/csv; filas nuevas que aportan: {sumadas}. Total final: {len(final)}')
    _resumen_cuentas(final, cortes)
    return final


# ------------------------------------------------------------------ 4b. mayor (fuente secundaria, opcional)
def mayor(cfg):
    """Suma, como hojas APARTE y rotuladas, los movimientos de un libro mayor contable para cuentas cuyo extracto
    del banco no existe (se perdio, no se descargo). Nunca se mezclan con las hojas del banco: el mayor no es
    un extracto y su saldo es relativo al periodo exportado. Se configura en `mayores` del proyecto.json."""
    t = _trabajo(cfg)
    filas_all, meta, excl_all = [], {}, []
    for m in cfg.get('mayores', []):
        ruta = m['archivo'] if os.path.isabs(m['archivo']) else os.path.join(cfg['entrada'], m['archivo'])
        desde = datetime.date.fromisoformat(m['desde']) if m.get('desde') else None
        hasta = datetime.date.fromisoformat(m['hasta']) if m.get('hasta') else None
        r = leer_mayor(ruta, desde, hasta)
        banco, cuenta = m['banco'], m['cuenta']
        filas = [dict(f, banco=banco, cuenta=cuenta, archivo=os.path.basename(ruta), nota='libro mayor') for f in r['filas']]
        filas_all.extend(filas)
        excl_all.extend(dict(f, banco=banco, cuenta=cuenta, archivo=os.path.basename(ruta), nota='fuera del periodo')
                        for f in r['excluidas'])
        meta[f'{banco}||{cuenta}'] = dict(solapa=m.get('solapa') or (cuenta + ' (mayor)'), archivo=os.path.basename(ruta),
                                          cuenta_mayor=r['cuenta_mayor'], excluidas=len(r['excluidas']))
        cortes = errores_de_saldo(filas)
        print(f'{banco} / {cuenta}: {len(filas)} filas del mayor ({r["cuenta_mayor"]}) '
              f'{filas[0]["fecha"] if filas else "-"} a {filas[-1]["fecha"] if filas else "-"}; '
              f'saldo corrido del mayor: {cortes} cortes; excluidas por periodo: {len(r["excluidas"])}')
        for f in r['excluidas']:
            print(f'   EXCLUIDA ({f["fecha"]}, {f["referencia"]}, {f["descripcion"][:40]}, '
                  f'debito {f["debito"]}, credito {f["credito"]}, cargada el {f["alta"]}): fuera del periodo pedido')
        for a in r['avisos']:
            print('   AVISO:', a)
    escribir_movimientos(os.path.join(t, 'mayor.csv'), filas_all)
    escribir_movimientos(os.path.join(t, 'mayor_excluidas.csv'), excl_all)
    json.dump(meta, open(os.path.join(t, 'mayor_meta.json'), 'w', encoding='utf-8'), ensure_ascii=False)
    if not cfg.get('mayores'):
        print('No hay `mayores` en la configuracion: nada que hacer.')
    return filas_all


# ------------------------------------------------------------------ 5. excel
ETIQUETAS = {
    'es': dict(titulo='Movimientos bancarios', sub='Fuente: extractos descargados del banco.', resumen='Resumen',
               cols=['Fecha', 'Descripción', 'Detalle', 'Referencia', 'Débito', 'Crédito', 'Saldo'],
               res=['Cuenta', 'Desde', 'Hasta', 'Movimientos', 'Saldo inicial', 'Total débitos', 'Total créditos', 'Saldo final'],
               notas='Notas', avisos='Avisos', av=['Cuenta', 'Fecha', 'Descripción', 'Saldo informado', 'Saldo de la fila anterior'],
               archivo='Movimientos', tz='Filas tomadas de otro archivo',
               mayor_titulo='Tomado del libro mayor contable (no hay extracto del banco para este tramo)',
               mayor_aviso='El mayor es el registro de la empresa, no lo informado por el banco. El saldo es acumulado desde el inicio de lo exportado y no es el saldo bancario.',
               mayor_saldo='Saldo acumulado del mayor', mayor_sec='Libro mayor (no es extracto del banco)'),
    'en': dict(titulo='Bank transactions', sub='Source: statements downloaded from the bank.', resumen='Summary',
               cols=['Date', 'Description', 'Details', 'Reference', 'Debit', 'Credit', 'Balance'],
               res=['Account', 'From', 'To', 'Transactions', 'Opening balance', 'Total debits', 'Total credits', 'Closing balance'],
               notas='Notes', avisos='Warnings', av=['Account', 'Date', 'Description', 'Reported balance', 'Previous row balance'],
               archivo='Transactions', tz='Rows taken from another file',
               mayor_titulo='Taken from the accounting ledger (no bank statement exists for this period)',
               mayor_aviso='The ledger is the company\'s own record, not what the bank reported. Its balance is cumulative from the start of the export and is not the bank balance.',
               mayor_saldo='Ledger running balance', mayor_sec='Accounting ledger (not a bank statement)'),
}


def _nombre_solapa(cuenta, usadas, solapas, banco):
    n = solapas.get(f'{banco}||{cuenta}') or cuenta
    n = re.sub(r'[\[\]\*\?/\\:]', ' ', n).strip()
    if len(n) > 31:
        n = n[-31:].strip()
    base, i = n, 2
    while n.lower() in usadas or n.lower() in ('resumen', 'summary', 'avisos', 'warnings'):
        n = (base[:28] + f' {i}')
        i += 1
    usadas.add(n.lower())
    return n


def excel(cfg):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    L = ETIQUETAS.get(cfg['idioma'], ETIQUETAS['es'])
    t = _trabajo(cfg)
    movs = leer_movimientos(os.path.join(t, 'final.csv'))
    cortes_all = list(__import__('csv').DictReader(open(os.path.join(t, 'cortes_final.csv'), encoding='utf-8-sig'), delimiter=';'))
    solapas = json.load(open(os.path.join(t, 'solapas.json'), encoding='utf-8')) if os.path.exists(os.path.join(t, 'solapas.json')) else {}
    notas_banco = cfg.get('notas', {})
    HEAD, BLANCO = PatternFill('solid', fgColor='1F3864'), Font(bold=True, color='FFFFFF')
    fmt = '#,##0.00'
    destino = cfg['salida']
    os.makedirs(os.path.join(destino, 'interno'), exist_ok=True)
    por_banco = collections.OrderedDict()
    for r in movs:
        por_banco.setdefault(r['banco'], collections.OrderedDict()).setdefault(r['cuenta'], []).append(r)
    may, may_meta = {}, {}
    if os.path.exists(os.path.join(t, 'mayor.csv')) and os.path.exists(os.path.join(t, 'mayor_meta.json')):
        may_meta = json.load(open(os.path.join(t, 'mayor_meta.json'), encoding='utf-8'))
        for r in leer_movimientos(os.path.join(t, 'mayor.csv')):
            may.setdefault(r['banco'], collections.OrderedDict()).setdefault(r['cuenta'], []).append(r)
        for b in may:
            por_banco.setdefault(b, collections.OrderedDict())
    traza, generados = [], []
    empresa = cfg['empresa'] + (f" (CUIT {cfg['cuit']})" if cfg['cuit'] else '')
    for banco, cuentas in por_banco.items():
        wb = Workbook()
        res = wb.active
        res.title = L['resumen']
        res['A1'] = f"{empresa} - {L['titulo']} {banco}"
        res['A1'].font = Font(bold=True, size=14)
        res['A2'] = (f"{cfg.get('periodo_texto', '')} {L['sub']}").strip()
        for c, x in enumerate(L['res'], 1):
            cel = res.cell(row=4, column=c, value=x); cel.fill, cel.font = HEAD, BLANCO
        usadas = set()
        for i, (cuenta, filas) in enumerate(cuentas.items()):
            ws = wb.create_sheet(_nombre_solapa(cuenta, usadas, solapas, banco))
            for c, x in enumerate(L['cols'], 1):
                cel = ws.cell(row=1, column=c, value=x); cel.fill, cel.font = HEAD, BLANCO
            for n, r in enumerate(filas, start=2):
                ws.cell(row=n, column=1, value=r['fecha']).number_format = 'DD/MM/YYYY'
                ws.cell(row=n, column=2, value=r['descripcion'])
                ws.cell(row=n, column=3, value=r['detalle'])
                ws.cell(row=n, column=4, value=r['referencia'])
                for c, k in ((5, 'debito'), (6, 'credito'), (7, 'saldo')):
                    ws.cell(row=n, column=c, value=None if r[k] is None else float(r[k])).number_format = fmt
                if r['nota']:
                    traza.append((banco, cuenta, r))
            for c, w in enumerate([12, 40, 50, 22, 16, 16, 18], 1):
                ws.column_dimensions[get_column_letter(c)].width = w
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = f'A1:G{len(filas) + 1}'
            p, u = filas[0], filas[-1]
            ini = (p['saldo'] + p['debito'] - p['credito']) if p['saldo'] is not None else None
            vals = [cuenta, p['fecha'], u['fecha'], len(filas), None if ini is None else float(ini),
                    float(sum(r['debito'] for r in filas)), float(sum(r['credito'] for r in filas)),
                    None if u['saldo'] is None else float(u['saldo'])]
            for c, v in enumerate(vals, 1):
                cel = res.cell(row=5 + i, column=c, value=v)
                if c in (2, 3): cel.number_format = 'DD/MM/YYYY'
                if c in (5, 6, 7, 8): cel.number_format = fmt
        for c, w in enumerate([34, 12, 12, 13, 18, 18, 18, 18], 1):
            res.column_dimensions[get_column_letter(c)].width = w
        res.freeze_panes = 'A5'
        fila = 6 + len(cuentas)
        if banco in may:
            res.cell(row=fila, column=1, value=L['mayor_sec']).font = Font(bold=True)
            for cuenta, filas in may[banco].items():
                fila += 1
                for c, v in enumerate([may_meta[f'{banco}||{cuenta}']['solapa'], filas[0]['fecha'], filas[-1]['fecha'], len(filas),
                                       None, float(sum(r['debito'] for r in filas)), float(sum(r['credito'] for r in filas)), None], 1):
                    cel = res.cell(row=fila, column=c, value=v)
                    if c in (2, 3): cel.number_format = 'DD/MM/YYYY'
                    if c in (6, 7): cel.number_format = fmt
            fila += 1
            res.cell(row=fila, column=1, value='• ' + L['mayor_aviso'])
            fila += 2
        notas = notas_banco.get(banco, [])
        if notas:
            res.cell(row=fila, column=1, value=L['notas']).font = Font(bold=True)
            for x in notas:
                fila += 1
                res.cell(row=fila, column=1, value='• ' + x)
        for cuenta, filas in (may.get(banco) or {}).items():      # hojas del mayor, aparte de las del banco
            meta = may_meta[f'{banco}||{cuenta}']
            wm = wb.create_sheet(_nombre_solapa(meta['solapa'], usadas, {}, banco))
            wm['A1'] = L['mayor_titulo']
            wm['A1'].font = Font(bold=True, size=12, color='9C0006')
            wm['A2'] = f"{L['mayor_aviso']} ({meta['archivo']}: {meta['cuenta_mayor']})"
            for c, x in enumerate(L['cols'][:-1] + [L['mayor_saldo']], 1):
                cel = wm.cell(row=4, column=c, value=x); cel.fill, cel.font = HEAD, BLANCO
            for n, r in enumerate(filas, start=5):
                wm.cell(row=n, column=1, value=r['fecha']).number_format = 'DD/MM/YYYY'
                wm.cell(row=n, column=2, value=r['descripcion'])
                wm.cell(row=n, column=3, value=r['detalle'])
                wm.cell(row=n, column=4, value=r['referencia'])
                for c, k in ((5, 'debito'), (6, 'credito'), (7, 'saldo')):
                    wm.cell(row=n, column=c, value=None if r[k] is None else float(r[k])).number_format = fmt
            for c, w in enumerate([12, 40, 60, 22, 16, 16, 22], 1):
                wm.column_dimensions[get_column_letter(c)].width = w
            wm.freeze_panes = 'A5'
            wm.auto_filter.ref = f'A4:G{len(filas) + 4}'
        cb = [k for k in cortes_all if k['banco'] == banco]
        if cb:      # los saldos que no empalman NUNCA se esconden
            av = wb.create_sheet(L['avisos'])
            for c, x in enumerate(L['av'], 1):
                cel = av.cell(row=1, column=c, value=x); cel.fill, cel.font = HEAD, BLANCO
            for n, k in enumerate(cb, start=2):
                av.cell(row=n, column=1, value=k['cuenta'])
                av.cell(row=n, column=2, value=datetime.date.fromisoformat(k['fecha'])).number_format = 'DD/MM/YYYY'
                av.cell(row=n, column=3, value=k['descripcion'])
                av.cell(row=n, column=4, value=float(k['saldo'])).number_format = fmt
                av.cell(row=n, column=5, value=float(k['saldo_anterior'])).number_format = fmt
        nombre = os.path.join(destino, f"{banco} {cfg['empresa']} - {L['archivo']}.xlsx")
        wb.save(nombre)
        generados.append((banco, nombre, sum(len(f) for f in cuentas.values()) + sum(len(f) for f in (may.get(banco) or {}).values())))
    wbt = Workbook()
    wt = wbt.active
    wt.title = L['tz'][:31]
    for c, x in enumerate(['Banco', 'Cuenta', 'Fecha', 'Descripción', 'Débito', 'Crédito', 'Saldo', 'Procedencia'], 1):
        cel = wt.cell(row=1, column=c, value=x); cel.fill, cel.font = HEAD, BLANCO
    for n, (b, cu, r) in enumerate(traza, start=2):
        for j, v in enumerate([b, cu, r['fecha'], r['descripcion'], float(r['debito']), float(r['credito']),
                               None if r['saldo'] is None else float(r['saldo']), r['nota']], 1):
            cel = wt.cell(row=n, column=j, value=v)
            if j == 3: cel.number_format = 'DD/MM/YYYY'
            if j in (5, 6, 7): cel.number_format = fmt
    for c, w in enumerate([14, 28, 12, 46, 16, 16, 18, 80], 1):
        wt.column_dimensions[get_column_letter(c)].width = w
    wt.freeze_panes = 'A2'
    wbt.save(os.path.join(destino, 'interno', 'Trazabilidad - filas tomadas de otro archivo.xlsx'))
    for b, nom, n in generados:
        print(f'{b:22} {n:5} filas -> {nom}')
    print(f'Trazabilidad interna: {len(traza)} filas')
    return generados


# ------------------------------------------------------------------ 6. verificar
def verificar(cfg):
    """Control independiente: reabre los Excel generados y recalcula todo. Devuelve 0 si todo cierra."""
    import openpyxl
    t = _trabajo(cfg)
    movs = leer_movimientos(os.path.join(t, 'final.csv'))
    L = ETIQUETAS.get(cfg['idioma'], ETIQUETAS['es'])
    problemas, total = [], 0
    por_banco = collections.defaultdict(list)
    for r in movs:
        por_banco[r['banco']].append(r)
    mayor_filas = collections.defaultdict(list)
    if os.path.exists(os.path.join(t, 'mayor.csv')):
        for r in leer_movimientos(os.path.join(t, 'mayor.csv')):
            mayor_filas[r['banco']].append(r)
            por_banco.setdefault(r['banco'], [])
    for banco, esperadas in por_banco.items():
        ruta = os.path.join(cfg['salida'], f"{banco} {cfg['empresa']} - {L['archivo']}.xlsx")
        if not os.path.exists(ruta):
            problemas.append(f'falta el Excel de {banco}')
            continue
        wb = openpyxl.load_workbook(ruta)
        filas_excel = 0
        suma_deb = suma_cred = CERO
        may_n, may_deb, may_cred = 0, CERO, CERO
        for ws in wb.worksheets:
            if ws.title in (L['resumen'], L['avisos']):
                continue
            es_mayor = ws['A1'].value == L['mayor_titulo']
            if es_mayor:
                for row in ws.iter_rows(min_row=5, values_only=True):
                    if row[0] is not None:
                        may_n += 1
                        may_deb += Decimal(str(row[4] or 0))
                        may_cred += Decimal(str(row[5] or 0))
                continue
            prev, cortes, n = None, 0, 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is None:
                    continue
                n += 1
                deb, cred = Decimal(str(row[4] or 0)), Decimal(str(row[5] or 0))
                sal = Decimal(str(row[6])) if row[6] is not None else None
                suma_deb += deb
                suma_cred += cred
                if prev is not None and sal is not None and abs(prev + cred - deb - sal) > TOL:
                    cortes += 1
                if sal is not None:
                    prev = sal
            filas_excel += n
            if cortes:
                problemas.append(f'{banco} / {ws.title}: {cortes} saldos que no empalman (recalculado desde el Excel)')
        if filas_excel != len(esperadas):
            problemas.append(f'{banco}: {filas_excel} filas en el Excel vs {len(esperadas)} esperadas')
        d = sum(r['debito'] for r in esperadas)
        c = sum(r['credito'] for r in esperadas)
        if abs(suma_deb - d) > TOL or abs(suma_cred - c) > TOL:
            problemas.append(f'{banco}: los totales de débito/crédito no coinciden con los datos')
        esperado_may = mayor_filas.get(banco, [])
        if may_n != len(esperado_may) or abs(may_deb - sum(r['debito'] for r in esperado_may)) > TOL \
                or abs(may_cred - sum(r['credito'] for r in esperado_may)) > TOL:
            problemas.append(f'{banco}: las hojas del mayor no coinciden con mayor.csv ({may_n} vs {len(esperado_may)} filas)')
        if may_n:
            print(f'  {banco:22} {may_n:5} filas del libro mayor en hojas aparte (rotuladas)')
        total += filas_excel
        print(f'  {banco:22} {filas_excel:5} filas en el Excel, totales {"OK" if abs(suma_deb - d) <= TOL and abs(suma_cred - c) <= TOL else "NO COINCIDEN"}')
    print(f'Total en Excel: {total} | datos finales: {len(movs)}')
    if problemas:
        print('PROBLEMAS:')
        for p in problemas:
            print('  -', p)
        return 1
    print('VERIFICACION OK: filas, totales y saldo corrido cierran en todas las cuentas.')
    return 0
